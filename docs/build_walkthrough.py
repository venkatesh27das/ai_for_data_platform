from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DOCS_DIR = Path(__file__).resolve().parent
ASSET_DIR = DOCS_DIR / "walkthrough-assets"
OUTPUT_PATH = DOCS_DIR / "Enterprise-Knowledge-Assembly-Studio-User-Walkthrough.docx"

INK = "111827"
SECONDARY = "4B5563"
MUTED = "6B7280"
PRIMARY = "F4511E"
PRIMARY_DARK = "D84315"
PRIMARY_SOFT = "FFF1EB"
BORDER = "E5E7EB"
SURFACE_MUTED = "F7F8FA"
SUCCESS = "159447"
BLUE = "2563EB"
WHITE = "FFFFFF"


SCREENS = [
    {
        "number": "1",
        "title": "Orient yourself on Home",
        "route": "/",
        "image": "01-home.jpg",
        "objective": "Use Home as the launch point for projects, recent work, platform activity, and enterprise-wide status.",
        "steps": [
            "Use the compact left navigation to move between Home, Knowledge Projects, Enterprise Assets, and Knowledge Products.",
            "Review the KPI cards for project, asset, product, quality, and agent-run health.",
            "Select New Knowledge Project to start an assembly workflow, or continue into the existing Customer 360 example.",
        ],
        "notice": "The top navigation stays generic. Customer 360 is the single implementation example, not the identity of the platform.",
        "alt": "Home page showing enterprise KPIs, recent Customer 360 work, activity, and the primary navigation.",
    },
    {
        "number": "2",
        "title": "Review the Knowledge Projects portfolio",
        "route": "/projects",
        "image": "02-projects.jpg",
        "objective": "Understand project status and enter either a new workflow or the existing Customer 360 workspace.",
        "steps": [
            "Review portfolio KPIs, lifecycle status, ownership, progress, and quality indicators.",
            "Open Customer 360 Knowledge Layer to follow the complete connected story.",
            "Select New Knowledge Project when you want to recreate the Customer 360 example through the guided wizard.",
        ],
        "notice": "A knowledge project is the governed assembly workspace. A knowledge product is the publishable result produced by that project.",
        "alt": "Knowledge Projects page showing the Customer 360 project and the New Knowledge Project action.",
    },
    {
        "number": "3",
        "title": "Explore governed Enterprise Assets",
        "route": "/assets",
        "image": "03-enterprise-assets.jpg",
        "objective": "Browse the governed assets that are available to be assembled into knowledge projects.",
        "steps": [
            "Filter by asset type, source system, domain, knowledge role, status, or quality.",
            "Select a source-system logo to narrow the inventory to Databricks, MDM, SharePoint, Atlan, dbt, or Neo4j.",
            "Review quality, owner, freshness, and readiness before including an asset in a project.",
        ],
        "notice": "This catalog represents outputs from existing enterprise platforms. The studio orchestrates them; it does not replace those systems.",
        "alt": "Enterprise Assets page with source-system logos, filters, KPI cards, and a governed asset inventory.",
    },
    {
        "number": "4",
        "title": "Inspect available Knowledge Products",
        "route": "/products",
        "image": "04-knowledge-products.jpg",
        "objective": "See which reusable knowledge products are published or in progress and who consumes them.",
        "steps": [
            "Select Customer 360 Knowledge Product from the table.",
            "Use the contextual inspector to review its version, owner, quality, endpoints, consumers, and governance state.",
            "Open the associated project when you need to change the assembly or prepare a new version.",
        ],
        "notice": "The prototype intentionally contains one knowledge product so the demonstration remains focused and internally consistent.",
        "alt": "Knowledge Products page showing the Customer 360 product and its contextual detail inspector.",
    },
    {
        "number": "5",
        "title": "Create a project: Scope & Domains",
        "route": "/projects/new",
        "image": "05-wizard-scope.jpg",
        "objective": "Define the business outcome, domain boundaries, target consumers, and expected knowledge product.",
        "steps": [
            "Confirm the project name and business objective, then choose the primary domain and sub-domains.",
            "Select target consumers such as RAG applications, copilots, AI agents, and analytics.",
            "Review the seeded business questions, sensitivity, residency, and approval controls before continuing.",
        ],
        "notice": "The wizard is pre-seeded with Customer 360 data so a conference demonstration can be completed quickly and deterministically.",
        "alt": "New Knowledge Project wizard on Scope and Domains with Customer 360 objective, domains, consumers, and business questions.",
    },
    {
        "number": "6",
        "title": "Create a project: Sources & Assets",
        "route": "/projects/new",
        "image": "06-wizard-sources.jpg",
        "objective": "Connect enterprise systems and select the governed assets that will contribute to the knowledge layer.",
        "steps": [
            "Verify that the required source systems are connected and healthy.",
            "Search and filter the discovered assets by type, domain, source, and knowledge role.",
            "Include or exclude assets and apply governed recommendations when they improve coverage.",
        ],
        "notice": "Asset selection is reversible. Recommendations are auditable suggestions and remain subject to human review.",
        "alt": "Sources and Assets wizard step showing connected platform logos and selectable enterprise asset rows.",
    },
    {
        "number": "7",
        "title": "Create a project: Success Criteria",
        "route": "/projects/new",
        "image": "07-wizard-success.jpg",
        "objective": "Translate the business objective into measurable quality, retrieval, coverage, and governance targets.",
        "steps": [
            "Review each success metric, its target, measurement method, and criticality.",
            "Adjust targets when the project has stricter domain or regulatory requirements.",
            "Confirm the acceptance criteria that must pass before the product can be published.",
        ],
        "notice": "Success criteria are used later by quality checks and publishing gates, keeping the project outcome measurable.",
        "alt": "Success Criteria wizard step showing target metrics and acceptance criteria for Customer 360.",
    },
    {
        "number": "8",
        "title": "Governance & Access",
        "route": "/projects/new",
        "image": "08-wizard-governance.jpg",
        "objective": "Attach classifications, policies, access groups, stewardship, and approval requirements before assembly begins.",
        "steps": [
            "Confirm information classifications and policy tags for customer and consent data.",
            "Review the selected governance policies and the access level granted to each enterprise group.",
            "Keep steward approval enabled for sensitive or PII-bearing knowledge products.",
        ],
        "notice": "Governance is configured before publication and remains visible throughout assembly, validation, and serving.",
        "alt": "Governance and Access wizard step showing policy controls, classifications, access groups, and approval settings.",
    },
    {
        "number": "9",
        "title": "Create a project: Review & Create",
        "route": "/projects/new",
        "image": "09-wizard-review.jpg",
        "objective": "Review the complete configuration and correct any issue before initializing the governed project.",
        "steps": [
            "Check project scope, selected sources, asset counts, success criteria, and governance controls.",
            "Use the Edit actions to return to a specific step without losing the rest of the configuration.",
            "Select Create Project when the review summary is complete.",
        ],
        "notice": "Creation starts deterministic agent services. The UI shows auditable objectives, evidence, confidence, and status rather than hidden model reasoning.",
        "alt": "Review and Create wizard step summarizing project scope, sources, selected assets, metrics, and governance.",
    },
    {
        "number": "10",
        "title": "Monitor project initialization",
        "route": "/projects/new",
        "image": "10-project-initialization.jpg",
        "objective": "Verify that the project was created and that governed discovery and assembly services have started.",
        "steps": [
            "Follow the initialization stages from project creation through asset discovery, semantic mapping, graph construction, validation, and draft publication.",
            "Review activated agents and the event log for current status and evidence.",
            "Open the project workspace once initialization is underway.",
        ],
        "notice": "The simulated services are deterministic for demo reliability, but the execution trace mirrors a governed enterprise-agent operating model.",
        "alt": "Project Created Successfully screen with initialization stages, agent status, project overview, and initial asset snapshot.",
    },
    {
        "number": "11",
        "title": "Read the Customer 360 project overview",
        "route": "/projects/customer-360",
        "image": "11-project-overview.jpg",
        "objective": "Use the overview as the shared status page for progress, health, activity, timeline, sources, and next governed actions.",
        "steps": [
            "Read the KPI row from Current Assembly through Data Freshness.",
            "Review project description, health, recent activity, timeline, and top source systems.",
            "Use Next Governed Actions or Quick Links to continue to the right workspace tab.",
        ],
        "notice": "The horizontal project tabs tell a connected story: connect, model, govern, validate, publish, and measure.",
        "alt": "Customer 360 project overview with assembly progress, project health, timeline, sources, and next actions.",
    },
    {
        "number": "12",
        "title": "Manage project Assets & Sources",
        "route": "/projects/customer-360/assets",
        "image": "12-assets-sources.jpg",
        "objective": "Confirm source health, review project-scoped assets, and improve coverage before graph assembly.",
        "steps": [
            "Review connected source systems, sync freshness, and asset counts.",
            "Search and filter the project inventory by source, domain, role, and status.",
            "Open recommendations when the coverage panel identifies high-impact missing assets.",
        ],
        "notice": "This view is project-scoped. The global Enterprise Assets page remains the broader inventory across the organization.",
        "alt": "Customer 360 Assets and Sources tab showing connected systems, project asset inventory, and coverage recommendations.",
    },
    {
        "number": "13",
        "title": "Explore the Graph & Model",
        "route": "/projects/customer-360/graph",
        "image": "13-graph-model.jpg",
        "objective": "Inspect the assembled knowledge graph, relationships, model layers, provenance, and downstream impact.",
        "steps": [
            "Use search, zoom, pan, fit-to-screen, layout, and layer controls to navigate the graph.",
            "Select a node or relationship to open its inspector and review provenance, confidence, policies, lineage, and connected assets.",
            "Use neighborhood expansion or shortest-path exploration to investigate how enterprise concepts are connected.",
        ],
        "notice": "The graph is a multi-layer enterprise knowledge representation, not only a diagram of tables or a database topology.",
        "alt": "Customer 360 Graph and Model tab showing a professional node-link graph, controls, legend, minimap, and right-side inspector.",
    },
    {
        "number": "14",
        "title": "Run Build & Govern",
        "route": "/projects/customer-360/build",
        "image": "14-build-govern.jpg",
        "objective": "Assemble semantic assets, attach governance controls, and review agent-generated recommendations.",
        "steps": [
            "Select Run Assembly to simulate the governed assembly pipeline.",
            "Follow concept extraction, entity alignment, semantic mapping, relationship inference, policy attachment, and provenance assembly.",
            "Review or reject recommendations and inspect agent runs when human judgment is required.",
        ],
        "notice": "Every agent action is presented as an auditable service execution with inputs, tools, output, evidence, confidence, duration, and review status.",
        "alt": "Build and Govern tab showing assembly progress, pipeline stages, governance readiness, and agent-run actions.",
    },
    {
        "number": "15",
        "title": "Validate Quality & Monitoring",
        "route": "/projects/customer-360/quality",
        "image": "15-quality-monitoring.jpg",
        "objective": "Measure whether the assembled knowledge layer satisfies quality, governance, freshness, and retrieval expectations.",
        "steps": [
            "Run a quality check and review the overall score and quality dimensions.",
            "Inspect warnings, failed checks, drift indicators, and the affected assets or relationships.",
            "Resolve critical findings before moving the draft into publication.",
        ],
        "notice": "Quality is evaluated against the success criteria configured during project creation, closing the loop between intent and evidence.",
        "alt": "Quality and Monitoring tab with overall score, dimension charts, quality checks, alerts, and monitoring details.",
    },
    {
        "number": "16",
        "title": "Publish & Serve the knowledge product",
        "route": "/projects/customer-360/publish",
        "image": "16-publish-serve.jpg",
        "objective": "Promote an approved assembly version into governed endpoints for applications, analytics, copilots, and agents.",
        "steps": [
            "Review version readiness, governance approvals, and the release pipeline.",
            "Inspect serving endpoints, access requirements, consumers, and deployment history.",
            "Publish the approved version or copy endpoint details for downstream integration.",
        ],
        "notice": "Publishing is governed promotion, not a data export. Lineage, policy bindings, ownership, and quality evidence remain attached.",
        "alt": "Publish and Serve tab showing release readiness, endpoints, consumers, deployment status, and publishing controls.",
    },
    {
        "number": "17",
        "title": "Measure Usage & Insights",
        "route": "/projects/customer-360/usage",
        "image": "17-usage-insights.jpg",
        "objective": "Understand adoption, endpoint usage, application demand, geography, search behavior, and improvement opportunities.",
        "steps": [
            "Review total consumption, active consumers, applications, API calls, response time, and satisfaction.",
            "Compare applications, knowledge assets, consumer types, and geographic distribution.",
            "Use search insights and recommendations to improve the next knowledge-product version.",
        ],
        "notice": "Usage completes the feedback loop: measured consumption and unmet queries become inputs to the next governed assembly cycle.",
        "alt": "Usage and Insights tab with consumption trends, top applications, geographic distribution, search insights, and recommendations.",
    },
    {
        "number": "18",
        "title": "Audit Activity and agent runs",
        "route": "/projects/customer-360/activity",
        "image": "18-activity.jpg",
        "objective": "Review what happened, which service performed it, what evidence was produced, and whether review is required.",
        "steps": [
            "Filter activity by type, status, agent, or time range.",
            "Open an agent run to inspect its objective, inputs, tools, output, evidence, confidence, duration, and human-review requirement.",
            "Use alerts and recent events to identify operational or governance follow-up.",
        ],
        "notice": "The activity model intentionally avoids conversational personas and hidden chain-of-thought. It exposes only governed execution evidence.",
        "alt": "Activity tab showing a filterable audit log, agent activity summary, alerts, and evidence-oriented run information.",
    },
    {
        "number": "19",
        "title": "Configure project Settings",
        "route": "/projects/customer-360/settings",
        "image": "19-settings.jpg",
        "objective": "Maintain project metadata, integrations, assembly behavior, access, notifications, and lifecycle controls.",
        "steps": [
            "Edit project identity, ownership, domain, and descriptive metadata.",
            "Review integrations, build settings, governance behavior, access rules, and notification preferences.",
            "Use lifecycle and danger-zone controls carefully because they change project operation or availability.",
        ],
        "notice": "Settings belong inside the project workspace, while top-level Administration remains reserved for platform-wide concerns.",
        "alt": "Customer 360 Settings tab showing project configuration, integrations, access, notifications, lifecycle, and administrative controls.",
    },
]


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_left_border(paragraph, color, size=18, space=6):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = p_bdr.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        p_bdr.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color)


def set_paragraph_bottom_border(paragraph, color, size=10, space=8):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)


def create_numbering(document, num_format="decimal", text="%1."):
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(element.get(qn("w:abstractNumId")))
        for element in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(element.get(qn("w:numId")))
        for element in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract.append(multi_level)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), num_format)
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    level.append(lvl_text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_numbered_item(document, text, num_id, keep_with_next=False):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = keep_with_next
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_element = OxmlElement("w:numId")
    num_id_element.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_element)
    p_pr.append(num_pr)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, color=INK)
    return paragraph


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = paragraph.add_run("Page ")
    set_run_font(label, size=8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = OxmlElement("w:r")
    run.append(begin)
    run.append(instruction)
    run.append(separate)
    run.append(text)
    run.append(end)
    paragraph._p.append(run)


def add_callout(document, label, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.06)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    set_paragraph_shading(paragraph, PRIMARY_SOFT)
    set_left_border(paragraph, PRIMARY)
    label_run = paragraph.add_run(f"{label}: ")
    set_run_font(label_run, size=9.5, color=PRIMARY_DARK, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=9.5, color=SECONDARY)
    return paragraph


def add_route(document, route):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.keep_with_next = True
    set_paragraph_shading(paragraph, SURFACE_MUTED)
    label = paragraph.add_run("Route  ")
    set_run_font(label, size=9, color=MUTED, bold=True)
    value = paragraph.add_run(route)
    set_run_font(value, name="Courier New", size=9, color=BLUE, bold=True)
    return paragraph


def add_screenshot(document, screen):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(2)
    picture = paragraph.add_run().add_picture(
        str(ASSET_DIR / screen["image"]),
        width=Inches(6.15),
    )
    picture._inline.docPr.set("descr", screen["alt"])
    picture._inline.docPr.set("title", screen["title"])

    caption = document.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    caption_run = caption.add_run(
        f"Figure {screen['number']}. {screen['title']} at 1600 x 1000 demo resolution."
    )
    set_run_font(caption_run, size=8.5, color=MUTED, italic=True)


def add_screen_page(document, screen):
    heading = document.add_paragraph(style="Heading 1")
    heading.paragraph_format.keep_with_next = True
    number_run = heading.add_run(f"{screen['number']}. ")
    set_run_font(number_run, size=16, color=PRIMARY_DARK, bold=True)
    title_run = heading.add_run(screen["title"])
    set_run_font(title_run, size=16, color=INK, bold=True)

    lead = document.add_paragraph(style="Guide Lead")
    lead.paragraph_format.keep_with_next = True
    lead_run = lead.add_run(screen["objective"])
    set_run_font(lead_run, size=10.8, color=SECONDARY)

    add_route(document, screen["route"])

    num_id = create_numbering(document)
    for index, item in enumerate(screen["steps"]):
        add_numbered_item(
            document,
            item,
            num_id,
            keep_with_next=index < len(screen["steps"]) - 1,
        )

    add_screenshot(document, screen)
    add_callout(document, "What to notice", screen["notice"])


def configure_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(INK)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.font.size = Pt(15)
    subtitle.font.color.rgb = RGBColor.from_string(SECONDARY)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(12)

    heading_values = {
        "Heading 1": (16, PRIMARY_DARK, 18, 10),
        "Heading 2": (13, PRIMARY_DARK, 14, 7),
        "Heading 3": (12, SECONDARY, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_values.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(5)
    caption.paragraph_format.line_spacing = 1.0

    if "Guide Lead" not in styles:
        guide_lead = styles.add_style("Guide Lead", WD_STYLE_TYPE.PARAGRAPH)
    else:
        guide_lead = styles["Guide Lead"]
    guide_lead.base_style = normal
    guide_lead.font.name = "Calibri"
    guide_lead._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    guide_lead._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    guide_lead.font.size = Pt(10.8)
    guide_lead.font.color.rgb = RGBColor.from_string(SECONDARY)
    guide_lead.paragraph_format.space_before = Pt(0)
    guide_lead.paragraph_format.space_after = Pt(5)
    guide_lead.paragraph_format.line_spacing = 1.15


def configure_page(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_paragraph.paragraph_format.space_after = Pt(0)
    header_run = header_paragraph.add_run(
        "ENTERPRISE KNOWLEDGE ASSEMBLY STUDIO  /  USER WALKTHROUGH"
    )
    set_run_font(header_run, size=8, color=MUTED, bold=True)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.paragraph_format.space_before = Pt(0)
    add_page_number(footer_paragraph)


def add_cover(document):
    for _ in range(4):
        document.add_paragraph()

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    kicker_run = kicker.add_run("USER WALKTHROUGH")
    set_run_font(kicker_run, size=11, color=PRIMARY, bold=True)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Enterprise Knowledge\nAssembly Studio")

    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "From governed enterprise assets to reusable Customer 360 knowledge products"
    )

    rule = document.add_paragraph()
    rule.paragraph_format.space_before = Pt(4)
    rule.paragraph_format.space_after = Pt(18)
    set_paragraph_bottom_border(rule, PRIMARY, size=20, space=6)

    summary = document.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary.paragraph_format.left_indent = Inches(0.55)
    summary.paragraph_format.right_indent = Inches(0.55)
    summary.paragraph_format.space_after = Pt(20)
    summary_run = summary.add_run(
        "A practical guide for creating, assembling, governing, validating, "
        "publishing, and measuring enterprise knowledge products."
    )
    set_run_font(summary_run, size=12, color=SECONDARY)

    metadata = document.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_after = Pt(5)
    metadata_run = metadata.add_run(
        "Prototype guide  |  Customer 360 implementation example  |  Version 1.0"
    )
    set_run_font(metadata_run, size=9.5, color=MUTED, bold=True)

    date_paragraph = document.add_paragraph()
    date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_run = date_paragraph.add_run("Updated July 29, 2026")
    set_run_font(date_run, size=9.5, color=MUTED)

    add_callout(
        document,
        "Audience",
        "Enterprise architects, data and AI platform leaders, governance teams, "
        "knowledge-product owners, and conference-demo facilitators.",
    )


def add_orientation_page(document):
    heading = document.add_paragraph(style="Heading 1")
    heading.add_run("How to use this guide")

    intro = document.add_paragraph()
    intro.add_run(
        "Follow the screens in sequence for a complete Customer 360 demonstration, "
        "or jump directly to a section when you need a specific task."
    )

    add_callout(
        document,
        "Start locally",
        "From the repository root, run ./start.sh and open "
        "http://127.0.0.1:4173. The prototype uses deterministic mock data and "
        "does not require API keys or an external backend.",
    )

    h2 = document.add_paragraph(style="Heading 2")
    h2.add_run("The connected story")
    story_steps = [
        "Discover governed assets produced by existing data, metadata, MDM, lineage, document, semantic, and graph platforms.",
        "Create a knowledge project with explicit business scope, quality targets, policies, access, and stewardship.",
        "Assemble entities, semantics, relationships, provenance, and policy bindings through governed agent services.",
        "Validate quality and retrieval readiness against measurable acceptance criteria.",
        "Publish approved knowledge products through endpoints for RAG, copilots, analytics, applications, and agents.",
        "Measure consumption and use activity evidence to improve the next governed version.",
    ]
    story_num_id = create_numbering(document)
    for item in story_steps:
        add_numbered_item(document, item, story_num_id)

    h2 = document.add_paragraph(style="Heading 2")
    h2.add_run("Navigation model")
    navigation_points = [
        ("Top-level navigation", "Home, Knowledge Projects, Enterprise Assets, Knowledge Products, and Administration."),
        ("Project tabs", "Overview, Assets & Sources, Graph & Model, Build & Govern, Quality & Monitoring, Publish & Serve, Usage & Insights, Activity, and Settings."),
        ("Contextual inspectors", "Open on the right when a selected asset, product, graph node, edge, recommendation, or run needs detail."),
    ]
    bullet_num_id = create_numbering(document, num_format="bullet", text="•")
    for label, detail in navigation_points:
        paragraph = add_numbered_item(document, "", bullet_num_id)
        paragraph.runs[0].text = ""
        label_run = paragraph.add_run(f"{label}: ")
        set_run_font(label_run, size=10.5, color=INK, bold=True)
        detail_run = paragraph.add_run(detail)
        set_run_font(detail_run, size=10.5, color=SECONDARY)

    add_callout(
        document,
        "Demo tip",
        "At 1600 x 1000 resolution, the full navigation, KPI rows, enterprise "
        "tables, and contextual inspectors remain visible without browser zoom.",
    )


def add_demo_script_page(document):
    heading = document.add_paragraph(style="Heading 1")
    heading.add_run("Suggested 12-minute demonstration")

    intro = document.add_paragraph()
    intro.add_run(
        "Use this sequence when presenting the product to an audience that needs "
        "the enterprise story before the implementation details."
    )

    demo_steps = [
        "Home (1 minute): establish the platform as a generic enterprise knowledge-assembly studio.",
        "Enterprise Assets (1 minute): show that governed platform outputs already exist and remain authoritative.",
        "New Project wizard (2 minutes): demonstrate scope, sources, metrics, governance, and review.",
        "Initialization (1 minute): explain governed agents and auditable execution evidence.",
        "Project Overview and Assets (1 minute): show connected project health and coverage.",
        "Graph & Model (2 minutes): search, select a node, inspect provenance, and explain graph layers.",
        "Build & Govern and Quality (2 minutes): run assembly, review recommendations, and validate readiness.",
        "Publish & Serve and Usage (2 minutes): show endpoints, consumers, adoption, query insights, and the feedback loop.",
    ]
    demo_num_id = create_numbering(document)
    for item in demo_steps:
        add_numbered_item(document, item, demo_num_id)

    h2 = document.add_paragraph(style="Heading 2")
    h2.add_run("Messages to reinforce")
    messages = [
        "The studio orchestrates existing enterprise capabilities; it does not replace the data platform, catalog, MDM, lineage, or document systems.",
        "Governance and evidence are attached throughout the workflow, not added after publication.",
        "Agents are governed enterprise services with observable objectives, tools, outputs, evidence, confidence, duration, and human-review requirements.",
        "Customer 360 is the implementation example. The platform and project-creation experience remain generic.",
    ]
    message_num_id = create_numbering(document, num_format="bullet", text="•")
    for item in messages:
        add_numbered_item(document, item, message_num_id)

    add_callout(
        document,
        "End state",
        "The audience should understand how distributed governed assets become a "
        "reusable, measurable, and publishable enterprise knowledge product.",
    )


def build_document():
    document = Document()
    configure_styles(document)
    configure_page(document)

    document.core_properties.title = (
        "Enterprise Knowledge Assembly Studio - User Walkthrough"
    )
    document.core_properties.subject = (
        "Step-by-step product walkthrough with Customer 360 screenshots"
    )
    document.core_properties.author = "Enterprise Knowledge Assembly Studio"
    document.core_properties.keywords = (
        "enterprise knowledge, Customer 360, knowledge graph, governance, walkthrough"
    )
    document.core_properties.comments = (
        "Generated from the local prototype using current application screenshots."
    )

    add_cover(document)
    document.add_page_break()
    add_orientation_page(document)

    for screen in SCREENS:
        document.add_page_break()
        add_screen_page(document, screen)

    document.add_page_break()
    add_demo_script_page(document)

    document.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    output = build_document()
    print(output)
