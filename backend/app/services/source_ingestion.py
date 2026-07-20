import csv
import io
import json
import re
from collections.abc import Iterable, Sequence
from datetime import date, datetime
from pathlib import Path
from typing import Any

from fastapi import HTTPException, UploadFile, status
from openpyxl import load_workbook

from app.db.models import ProjectSource
from app.repositories.sources import SourceRepository

TEXT_FORMATS = {"csv", "json", "ddl", "sql"}
SUPPORTED_FORMATS = TEXT_FORMATS | {"xlsx", "xlsm"}
PROFILE_ROWS = 100


class SourceIngestionService:
    def __init__(self, repository: SourceRepository, max_upload_mb: int) -> None:
        self.repository = repository
        self.max_bytes = max_upload_mb * 1024 * 1024

    def list(self, project_id: str) -> list[ProjectSource]:
        return self.repository.list_for_project(project_id)

    async def ingest(self, project_id: str, upload: UploadFile) -> ProjectSource:
        name = Path(upload.filename or "source").name
        source_format = Path(name).suffix.lower().lstrip(".")
        if source_format not in SUPPORTED_FORMATS:
            raise HTTPException(
                status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
                detail="Supported source files are CSV, JSON, DDL, SQL, XLSX, and XLSM.",
            )
        content = await upload.read(self.max_bytes + 1)
        if len(content) > self.max_bytes:
            raise HTTPException(
                status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                detail=f"Source file exceeds the {self.max_bytes // 1024 // 1024} MB limit.",
            )
        try:
            profile, excerpt = profile_source(name, source_format, content)
        except (UnicodeDecodeError, csv.Error, json.JSONDecodeError, ValueError) as exc:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Could not inspect {name}: {exc}",
            ) from exc
        return self.repository.create(
            project_id=project_id,
            name=name,
            format=source_format,
            size_bytes=len(content),
            content_excerpt=excerpt,
            profile=profile,
        )


def profile_source(name: str, source_format: str, content: bytes) -> tuple[dict[str, Any], str]:
    if source_format == "csv":
        text = content.decode("utf-8-sig")
        return {"tables": [profile_csv(Path(name).stem, text)]}, text[:4_000]
    if source_format == "json":
        text = content.decode("utf-8-sig")
        return {"tables": profile_json(Path(name).stem, json.loads(text))}, text[:4_000]
    if source_format in {"ddl", "sql"}:
        text = content.decode("utf-8-sig")
        return {"tables": profile_ddl(text)}, text[:20_000]
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    tables = []
    try:
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(max_row=PROFILE_ROWS + 1, values_only=True))
            if rows:
                tables.append(profile_rows(sheet.title, rows[0], rows[1:]))
    finally:
        workbook.close()
    return {"tables": tables}, json.dumps({"workbook": name, "sheets": tables})[:20_000]


def profile_csv(name: str, text: str) -> dict[str, Any]:
    sample = text[:8192]
    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    reader = csv.reader(io.StringIO(text), dialect)
    rows = list(_take(reader, PROFILE_ROWS + 1))
    if not rows:
        return profile_rows(name, [], [])
    return profile_rows(name, rows[0], rows[1:])


def profile_json(name: str, value: Any) -> list[dict[str, Any]]:
    if isinstance(value, list):
        records = [item for item in value[:PROFILE_ROWS] if isinstance(item, dict)]
        headers = ordered_keys(records)
        return [profile_rows(name, headers, [[row.get(key) for key in headers] for row in records])]
    if isinstance(value, dict):
        nested = [
            profile_json(str(key), item)[0]
            for key, item in value.items()
            if isinstance(item, list) and item and isinstance(item[0], dict)
        ]
        if nested:
            return nested
        return [profile_rows(name, list(value), [list(value.values())])]
    raise ValueError("JSON must contain an object or an array of objects")


def profile_ddl(text: str) -> list[dict[str, Any]]:
    tables: list[dict[str, Any]] = []
    for match in re.finditer(
        r"CREATE\s+TABLE\s+(?P<name>[A-Za-z0-9_.\"`\[\]]+)\s*\((?P<body>.*?)\)\s*;",
        text,
        re.IGNORECASE | re.DOTALL,
    ):
        columns: list[dict[str, Any]] = []
        for definition in split_ddl_definitions(match.group("body")):
            if definition.upper().startswith(
                ("PRIMARY KEY", "FOREIGN KEY", "CONSTRAINT", "UNIQUE", "CHECK")
            ):
                continue
            parts = definition.split()
            if parts:
                columns.append(
                    {
                        "name": parts[0].strip('"`[]'),
                        "data_type": parts[1] if len(parts) > 1 else "unknown",
                        "null_count": None,
                        "sample_values": [],
                    }
                )
        tables.append(
            {
                "name": match.group("name").strip('"`[]').split(".")[-1],
                "row_sample_count": 0,
                "columns": columns,
            }
        )
    return tables


def profile_rows(
    name: str, headers: Iterable[Any], rows: Sequence[Iterable[Any]]
) -> dict[str, Any]:
    names = [
        str(value).strip() if value is not None else f"column_{index + 1}"
        for index, value in enumerate(headers)
    ]
    materialized = [list(row) for row in rows]
    columns = []
    for index, column_name in enumerate(names):
        values = [row[index] if index < len(row) else None for row in materialized]
        populated = [value for value in values if value not in (None, "")]
        columns.append(
            {
                "name": column_name,
                "data_type": infer_type(populated),
                "null_count": len(values) - len(populated),
                "sample_values": [serialize_value(value) for value in populated[:3]],
            }
        )
    return {"name": name, "row_sample_count": len(materialized), "columns": columns}


def infer_type(values: list[Any]) -> str:
    if not values:
        return "unknown"
    if all(isinstance(value, bool) for value in values):
        return "boolean"
    if all(isinstance(value, int) and not isinstance(value, bool) for value in values):
        return "integer"
    if all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in values):
        return "number"
    if all(isinstance(value, (date, datetime)) for value in values):
        return "date"
    if all(isinstance(value, str) and re.fullmatch(r"[-+]?\d+", value.strip()) for value in values):
        return "integer"
    if all(
        isinstance(value, str) and re.fullmatch(r"[-+]?(?:\d+\.?\d*|\.\d+)", value.strip())
        for value in values
    ):
        return "number"
    if all(
        isinstance(value, str)
        and re.fullmatch(r"\d{4}-\d{2}-\d{2}(?:[T ][0-9:.+\-Z]+)?", value.strip())
        for value in values
    ):
        return "date"
    return "string"


def serialize_value(value: Any) -> str | int | float | bool:
    if isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def ordered_keys(records: list[dict[str, Any]]) -> list[str]:
    return list(dict.fromkeys(key for record in records for key in record))


def _take(rows: Iterable[list[str]], count: int) -> Iterable[list[str]]:
    for index, row in enumerate(rows):
        if index >= count:
            break
        yield row


def split_ddl_definitions(body: str) -> list[str]:
    result: list[str] = []
    current: list[str] = []
    depth = 0
    for character in body:
        depth += 1 if character == "(" else -1 if character == ")" else 0
        if character == "," and depth == 0:
            result.append("".join(current).strip())
            current = []
        else:
            current.append(character)
    if current:
        result.append("".join(current).strip())
    return result
